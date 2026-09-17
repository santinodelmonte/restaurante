#!/usr/bin/env python3
"""Arma la planilla de clasificación que se le manda a Bo Fomento.

Lee el MENU de index.html y escribe docs/planilla-clasificacion-bo-fomento.xlsx
con los platos ya cargados y nuestra lectura preliminar de la carta impresa.
Lo que el local corrija en la planilla vuelve después al array MENU.

Uso:  python3 docs/planilla.py      (requiere openpyxl)
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "index.html"
SALIDA = RAIZ / "docs" / "planilla-clasificacion-bo-fomento.xlsx"

MARINO, ESPUMA, ONDA, ARENA = "1B3A57", "F6F9FB", "CBD9E3", "FDF0D8"

CATS = {
    "entradas": "Entradas", "horno": "Horno de barro", "parrilla": "Parrilla a toda brasa",
    "cervezas": "Cervezas", "tragos": "Tragos", "refrescos": "Refrescos y licuados",
    "bebidas": "Bebidas y whiskys", "postres": "Postres", "cafeteria": "Cafetería",
}

ALERGENOS = [
    ("gluten", "Gluten"), ("lacteos", "Lácteos"), ("huevo", "Huevo"),
    ("pescado", "Pescado"), ("mariscos", "Mariscos"), ("frutos-secos", "Frutos secos"),
    ("cerdo", "Cerdo"), ("carne", "Carne"), ("pollo", "Pollo"), ("alcohol", "Alcohol"),
]

COLUMNAS = [
    ("#", 5), ("Categoría", 20), ("Sección", 18), ("Plato", 30), ("Precio $", 10),
    ("¿Sigue en carta?", 15), ("Descripción para la carta", 46),
    ("Vegetariano", 12), ("Vegano", 11), ("Sin gluten", 12), ("Para compartir", 14),
] + [(n, 11) for _, n in ALERGENOS] + [("Picante", 10), ("Notas / correcciones", 34)]


def leer_menu():
    """Saca el array MENU de index.html sin ejecutar JavaScript."""
    src = FUENTE.read_text(encoding="utf-8")
    i = src.index("const MENU = [")
    inicio = src.index("[", i)
    prof, j = 0, inicio
    while j < len(src):
        if src[j] == "[":
            prof += 1
        elif src[j] == "]":
            prof -= 1
            if prof == 0:
                j += 1
                break
        j += 1
    crudo = src[inicio:j]
    crudo = re.sub(r"/\*.*?\*/", "", crudo, flags=re.S)       # comentarios de sección
    crudo = re.sub(r"(?m)//.*$", "", crudo)
    crudo = re.sub(r"([{,]\s*)([A-Za-z_][A-Za-z0-9_]*)\s*:", r'\1"\2":', crudo)  # claves con comillas
    crudo = re.sub(r",(\s*[}\]])", r"\1", crudo)              # comas colgadas
    return json.loads(crudo)


def si_no(v):
    return "sí" if v else "no"


def es_vegano(p):
    """Propuesta preliminar: vegetariano y sin ningún ingrediente animal declarado."""
    if not p["apto"].get("vegetariano"):
        return False
    return not ({"lacteos", "huevo", "pescado", "mariscos", "carne", "cerdo", "pollo", "achuras"}
                & set(p.get("contiene", [])))


def encabezado(ws, fila=1):
    borde = Side(style="thin", color=MARINO)
    for col, (titulo, ancho) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=fila, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=MARINO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=borde)
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.row_dimensions[fila].height = 32
    ws.freeze_panes = f"E{fila + 1}"
    ws.auto_filter.ref = f"A{fila}:{get_column_letter(len(COLUMNAS))}{fila}"


def validaciones(ws, filas):
    """Desplegables sí/no en todas las columnas que se marcan."""
    dv = DataValidation(type="list", formula1='"sí,no"', allow_blank=True)
    dv.error = "Poné sí o no."
    dv.errorTitle = "Valor no válido"
    ws.add_data_validation(dv)
    primera, ultima = 2, max(filas, 2)
    for col in range(6, len(COLUMNAS)):          # de "¿Sigue en carta?" a "Picante"
        if col == 7:                              # la descripción es texto libre
            continue
        letra = get_column_letter(col)
        dv.add(f"{letra}{primera}:{letra}{ultima}")


def hoja_instrucciones(wb, total):
    ws = wb.create_sheet("Cómo completar", 0)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 104
    texto = [
        ("t", "PLANILLA DE CLASIFICACIÓN — BO FOMENTO"),
        ("s", "Carta digital + el genio · Aura Marketing · Setiembre 2026"),
        ("", ""),
        ("h", "Qué hay que hacer"),
        ("p", f"En la hoja «Platos» están los {total} platos que ya cargamos, con el precio que tomamos "
              "de la carta impresa. Cada fila se marca con sí o no en las columnas de color."),
        ("p", "Las casillas ya vienen con NUESTRA LECTURA de la carta. Es una propuesta, no un dato "
              "confirmado: lo que no corrijan, lo damos por bueno y sale publicado así."),
        ("p", "En la hoja «Faltan» van los vinos y las milanesas / burgers, que no estaban en las "
              "imágenes que recibimos. Se agregan ahí con el mismo formato."),
        ("", ""),
        ("h", "Qué significa cada columna"),
        ("p", "Vegetariano · sin carne, pollo, pescado ni mariscos. Puede llevar lácteos y huevo."),
        ("p", "Vegano · además, sin lácteos, huevo ni miel. Ojo con la manteca de la plancha y los aderezos."),
        ("p", "Sin gluten · sin trigo, avena, cebada ni centeno: harina, pan rallado, rebozados, "
              "salsas espesadas, cerveza."),
        ("p", "Para compartir · si tiene sentido pedirlo para el centro de la mesa."),
        ("p", "Alérgenos · marcar sí en todo lo que el plato contenga, aunque sea en poca cantidad."),
        ("p", "Alcohol · para que el genio no se lo recomiende a quien pidió algo sin alcohol."),
        ("p", "Picante · para avisarlo en la ficha del plato."),
        ("p", "¿Sigue en carta? · poner no en los platos que ya no se hacen. Esos no se publican."),
        ("p", "Descripción · una línea por plato: qué trae y con qué viene. Es lo que más vende y lo que "
              "el genio usa para justificar la recomendación."),
        ("", ""),
        ("h", "Lo importante"),
        ("p", "En el sitio, los filtros de vegetariano, vegano y sin gluten SACAN platos de la "
              "recomendación: no los penalizan, los eliminan. Un plato mal marcado se le recomienda a la "
              "persona equivocada con total seguridad. Por eso pedimos que esta planilla la revise alguien "
              "de cocina antes de devolverla."),
        ("p", "Contaminación cruzada: si un plato no lleva harina pero se fríe en el mismo aceite que los "
              "rebozados, no es apto para celíacos. En ese caso va no en «Sin gluten» y se aclara en Notas."),
        ("p", "Publicamos exactamente lo que ustedes marcan y no lo cambiamos por nuestra cuenta."),
        ("", ""),
        ("h", "Cuando esté pronta"),
        ("p", "Nos la mandan por WhatsApp o mail junto con la carta completa y el logo. Ese es el día 0: "
              "desde ahí son 10 a 12 días hábiles hasta publicarlo."),
        ("", ""),
        ("s", "Aura Marketing · +598 92 681 102 · auramkt.uy"),
    ]
    fila = 2
    for tipo, linea in texto:
        c = ws.cell(row=fila, column=2, value=linea)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if tipo == "t":
            c.font = Font(bold=True, size=16, color=MARINO)
        elif tipo == "s":
            c.font = Font(size=10, color="4A6379")
        elif tipo == "h":
            c.font = Font(bold=True, size=12, color=MARINO)
        else:
            c.font = Font(size=10)
            ws.row_dimensions[fila].height = 15 * (1 + len(linea) // 95)
        fila += 1
    return ws


def main():
    menu = leer_menu()
    wb = Workbook()
    ws = wb.active
    ws.title = "Platos"
    encabezado(ws)

    marca = PatternFill("solid", fgColor=ARENA)
    raya = PatternFill("solid", fgColor=ESPUMA)
    borde = Border(bottom=Side(style="hair", color=ONDA))

    for n, p in enumerate(menu, start=1):
        contiene = set(p.get("contiene", []))
        fila = [
            n, CATS.get(p["categoria"], p["categoria"]), p.get("sub", ""), p["nombre"],
            p["precio"], "sí", p.get("descripcion", ""),
            si_no(p["apto"].get("vegetariano")), si_no(es_vegano(p)),
            si_no(p["apto"].get("singluten")), si_no(p["apto"].get("compartir")),
        ] + [si_no(k in contiene) for k, _ in ALERGENOS] + ["no", ""]

        r = n + 1
        for col, valor in enumerate(fila, start=1):
            c = ws.cell(row=r, column=col, value=valor)
            c.font = Font(size=9)
            c.border = borde
            c.alignment = Alignment(
                horizontal="center" if col >= 6 and col != 7 else "left",
                vertical="top", wrap_text=col in (4, 7, len(COLUMNAS)),
            )
            if n % 2 == 0 and col < 6:
                c.fill = raya
        # Lo que declara «sin gluten» queda resaltado: es lo más riesgoso de la planilla.
        if p["apto"].get("singluten"):
            ws.cell(row=r, column=10).fill = marca
        if not p.get("descripcion"):
            ws.cell(row=r, column=7).fill = marca

    validaciones(ws, len(menu) + 1)

    # ── Hoja de lo que falta ────────────────────────────────────────
    wf = wb.create_sheet("Faltan")
    encabezado(wf)
    pendientes = [("Vinos", 40), ("Milanesas y burgers", 25)]
    fila = 2
    for seccion, cuantos in pendientes:
        for _ in range(cuantos):
            wf.cell(row=fila, column=2, value=seccion).font = Font(size=9)
            for col in range(1, len(COLUMNAS) + 1):
                wf.cell(row=fila, column=col).border = borde
            fila += 1
    validaciones(wf, fila)

    hoja_instrucciones(wb, len(menu))
    wb.active = 0
    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Planilla lista: {SALIDA} ({len(menu)} platos)")


if __name__ == "__main__":
    main()
